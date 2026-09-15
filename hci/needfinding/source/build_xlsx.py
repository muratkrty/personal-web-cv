"""Build iCanvas_needfinding.xlsx from questions.json + out/responses_final.csv + out/analysis.json"""
import json, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.chart import BarChart, LineChart, ScatterChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.comments import Comment

Q = json.load(open("questions.json"))
QI = {q["id"]: q for s in Q["sections"] for q in s["questions"]}
A = json.load(open("out/analysis.json"))
rows = list(csv.DictReader(open("out/responses_final.csv")))
N = len(rows)
LAST = 500          # formulas cover rows 2..500 so real exports can be appended
SEGS = A["segment_order"]

# ----------------------------------------------------------------- styles
FONT = "Arial"
def f(size=10, bold=False, color="222222", italic=False): return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
def fill(hex_): return PatternFill("solid", start_color=hex_, end_color=hex_)
TEAL, TEAL_D, GREEN = "1B7F5E", "145F46", "1B7F5E"
LAV, LAV_D, PINK, GREY, GREY_D = "DCDAFB", "C9C6F5", "FFF4FA", "F3F4F6", "E5E7EB"
thin = Side(style="thin", color="D1D5DB")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

def title(ws, cell, text, size=16):
    ws[cell] = text; ws[cell].font = f(size, True, TEAL_D)
def note(ws, cell, text, italic=True, color="555555"):
    ws[cell] = text; ws[cell].font = f(9, False, color, italic); ws[cell].alignment = Alignment(wrap_text=True, vertical="top")
def header_row(ws, r, c0, headers, fill_hex=TEAL, color="FFFFFF"):
    for i, h in enumerate(headers):
        c = ws.cell(row=r, column=c0 + i, value=h); c.font = f(10, True, color); c.fill = fill(fill_hex); c.alignment = CENTER; c.border = BORDER
def cellw(ws, r, c, v, bold=False, fmt=None, fill_hex=None, align=WRAP, color="222222", size=10):
    x = ws.cell(row=r, column=c, value=v); x.font = f(size, bold, color); x.alignment = align; x.border = BORDER
    if fmt: x.number_format = fmt
    if fill_hex: x.fill = fill(fill_hex)
    return x
def opt_label(qid, k): return next(o["label"] for o in QI[qid]["options"] if o["k"] == k)
def section_title(ws, r, text, cols=12):
    ws.cell(row=r, column=1, value=text).font = f(12, True, "FFFFFF")
    for c in range(1, cols + 1): ws.cell(row=r, column=c).fill = fill(TEAL_D)
    ws.row_dimensions[r].height = 20

wb = Workbook()

# =================================================================== README
ws = wb.active; ws.title = "README"
ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 28; ws.column_dimensions["C"].width = 62; ws.column_dimensions["D"].width = 40
title(ws, "B2", "iCanvas needfinding: survey data, patterns, segments and a persona", 16)
note(ws, "B3", "Human-Computer Interaction (Tilburg University, AIR-Lab). Companion to the Week 02 (UCD & needfinding) and Week 03 (design alternatives) lectures.", False)
r = 5
section_title(ws, r, "What this workbook is", 4); r += 1
txt = (f"A design team is proposing iCanvas, a new lecture management system for Tilburg University. Before ideating, the team ran an "
       f"anonymous online survey (needfinding) with {N} respondents about how they manage their courses in the current system (Canvas). "
       "This workbook contains the questionnaire, the raw responses and the three analysis steps that turn responses into a persona "
       "(Cooper et al., About Face): 1. analyse data to find patterns, 2. cluster user segments, 3. write a persona based on a segment.")
ws.cell(row=r, column=2, value=txt).alignment = WRAP; ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); ws.row_dimensions[r].height = 75; ws.cell(row=r, column=2).font = f(10)
r += 2
section_title(ws, r, "Stakeholders (Dix et al.) and who is in the sample", 4); r += 1
header_row(ws, r, 2, ["Stakeholder type", "Who", "In this survey?"]); r += 1
for a, b, c in [("Primary (use iCanvas)", "Students (BSc, MSc, PhD), lecturers/professors, teaching assistants", f"Yes: {sum(1 for x in rows if x['A1_role'] in ('bsc','msc','phd'))} students, {sum(1 for x in rows if x['A1_role'] in ('lecturer','ta'))} lecturers/TAs"),
                ("Secondary (provide input / receive output)", "Education office, IT support, library staff; parents; education director", f"Partly: {sum(1 for x in rows if x['A1_role']=='staff')} support/admin staff"),
                ("Tertiary (affected, no direct use)", "Department heads, examination board", "No (interviews planned)"),
                ("Facilitating (build / deploy)", "Design team, TiU IT department", "No (they are us)"),
                ("Extreme users (slide: Participants)", "Screen-reader / caption / dyslexia users; lecturers who use Canvas all day vs. almost never; students with 24+ h/week of work or care", f"Yes, deliberately recruited: {sum(1 for x in rows if any(x[f'A8_access__{k}']=='1' for k in ['screenreader','captions','dyslexia','motor','contrast']))} with accessibility needs, {sum(1 for x in rows if x['A7_hours'] in ('17-24','25+'))} with 17+ h outside commitments")]:
    cellw(ws, r, 2, a, True); cellw(ws, r, 3, b); cellw(ws, r, 4, c); ws.row_dimensions[r].height = 45; r += 1
r += 1
section_title(ws, r, "Method", 4); r += 1
for a, b in [("Technique", "Online survey (questionnaire), the needfinding technique highlighted in green on the Week 02 slides. Interviews and observation are the planned follow-up for the surprising findings."),
             ("Instrument", "34 items in 5 sections: A About you (demographics, Muller et al. Table 55.1 style), B Habits (frequency, devices, features, deadline discovery), C Experience (10 Likert statements incl. one inverted duplicate pair, 7-point satisfaction, 5-point difficulty), D Ranking of 6 activities, E three open-ended grand-tour questions. See sheet 'Questionnaire'."),
             ("Question rules applied", "No leading questions ('Is the calendar important to you?'), no 'what would you like in a tool?', no hypotheticals; ask about the last concrete time something happened (grand tour). Close-ended where the universe of answers is known and small; open-ended otherwise. Age asked in ranges and no identifiers collected (privacy)."),
             ("Sampling", "Convenience sample through CSAI/DSS mailing lists, the Mensa and the education office, topped up with deliberately recruited extreme users. Respondents are not a random sample of TiU; treat percentages as indicative."),
             ("Period", "8-14 September 2026 (see column 'submitted')."),
             ("Analysis", "Descriptive statistics and cross-tabulations (sheet 1-Patterns); k-means clustering on standardised behaviour + attitude features with k chosen by elbow/silhouette and interpretability (sheet 2-Segments); one persona written from the largest segment and short persona candidates for the others (sheet 3-Persona). All statistics on sheets 1 and 2 are live formulas on the Responses sheet; cluster labels and PCA coordinates are pasted values because Excel cannot run k-means.")]:
    cellw(ws, r, 2, a, True); cellw(ws, r, 3, b); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4); ws.row_dimensions[r].height = 62; r += 1
r += 1
section_title(ws, r, "Sheet map (follows the persona process on the Week 03 slides)", 4); r += 1
header_row(ws, r, 2, ["Sheet", "Contents", "Persona process step"]); r += 1
for a, b, c in [("Questionnaire", "Every question, its type, answer options, numeric coding and the design rationale.", "Instrument"),
                ("Codes", "Lookup table code -> label -> numeric value, used by the formulas.", "Coding"),
                ("Responses", f"{N} simulated respondents, one per row. Codes for single-choice items, 1/0 columns for multi-select items, rank 1-6 per activity, verbatim text. Derived numeric columns and the cluster label are at the right. vis.html reads this sheet.", "1. Collect user needs/data"),
                ("1-Patterns", "Sample description, habits, attitudes (means, SD, distributions), reliability of the inverted pair, rankings, cross-tabs, correlations, key patterns.", "2. Analyse data to find patterns"),
                ("2-Segments", "k selection, segment sizes, profile matrix per segment, role x segment, PCA map, segment descriptions with quotes.", "3. Cluster user segments"),
                ("3-Persona", "Full persona card for the largest segment (template layout), evidence table linking every claim to a statistic, POV statement, How-Might-We questions, persona candidates for the other segments.", "4. Write a persona based on a segment")]:
    cellw(ws, r, 2, a, True); cellw(ws, r, 3, b); cellw(ws, r, 4, c); ws.row_dimensions[r].height = 42; r += 1
r += 1
section_title(ws, r, "Coding conventions", 4); r += 1
for a, b in [("Single choice", "Stored as the option code (e.g. A1_role = 'bsc'). Sheet 'Codes' maps codes to labels."),
             ("Ordinal scales", "Frequency B1 (1 = less than once a week ... 6 = several times a day), visit length B4 (1-5), outside commitments A7 (0-4), times lost B8 (0-4) follow the numerical equivalences of Muller et al. (Table 55.2). Derived *_v columns hold the numbers."),
             ("Likert C1-C10", "1 = Strongly disagree ... 7 = Strongly agree. C2 is the inverted duplicate of C1; reverse it (8 - C2) before averaging with C1. C11 satisfaction 1-7, C12 difficulty 1-5 (1 = very difficult)."),
             ("Multi-select", "One column per option, 1 = selected, 0 = not selected (A8, B2, B5, B6, B9)."),
             ("Ranking D1", "One column per activity holding its rank, 1 = most important ... 6 = least important."),
             ("Open text E1-E3", "Verbatim. Empty = skipped.")]:
    cellw(ws, r, 2, a, True); cellw(ws, r, 3, b); ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=4); ws.row_dimensions[r].height = 34; r += 1
r += 1
section_title(ws, r, "Data provenance (read this)", 4); r += 1
prov = ("The responses in this workbook are SIMULATED for teaching. They were generated (seed 2026) from six latent behavioural profiles of "
        "Canvas users at a Dutch university, with realistic noise, a straight-liner and skipped optional items, so that the three analysis "
        "steps can be demonstrated end to end. Nothing here is a statement about real Tilburg University users. The companion page vis.html "
        "reads this Responses sheet when it opens, replays the simulated submissions one by one and runs the same three analysis steps in the browser.")
ws.cell(row=r, column=2, value=prov).alignment = WRAP; ws.cell(row=r, column=2).font = f(10, False, "7A1F1F"); ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); ws.row_dimensions[r].height = 78
r += 2
note(ws, f"B{r}", "Companion files: vis.html (reads the Responses sheet of this workbook and runs the same analysis in the browser) and user.html (the same questionnaire as a stand-alone web page for students during the lecture).", False)

# =================================================================== Questionnaire
ws = wb.create_sheet("Questionnaire")
widths = {"A": 6, "B": 10, "C": 52, "D": 12, "E": 46, "F": 26, "G": 30, "H": 46}
for k, v in widths.items(): ws.column_dimensions[k].width = v
title(ws, "A1", "Questionnaire and codebook", 14)
note(ws, "A2", Q["intro"] + "  |  Privacy notice shown to respondents: " + Q["privacy"], False); ws.merge_cells("A2:H2"); ws.row_dimensions[2].height = 48
header_row(ws, 4, 1, ["ID", "Section", "Question", "Type", "Answer options", "Numeric coding", "Column(s) in Responses", "Design rationale (lecture slides)"])
RATIONALE = {
 "A1": "Role first, so needs can be compared across primary stakeholders (students vs. lecturers vs. staff).",
 "A2": "Age in ranges rather than exact age: enough for segmentation, better for privacy (slide: consider privacy of users).",
 "A3": "Inclusive options plus 'prefer not to say'; optional.",
 "A4": "School as a proxy for discipline/culture of teaching.",
 "A5": "Culture proxy that is less intrusive than nationality (slide: age, culture, abilities, technology familiarity...).",
 "A6": "Load: how many courses compete for attention.",
 "A7": "Identifies extreme users with heavy outside commitments (jobs, caring). Ordinal, coded 0-4.",
 "A8": "Abilities/disabilities: recruits the accessibility extreme users; multi-select with exclusive 'none' and 'prefer not to say'.",
 "A9": "Technology familiarity, 7-point (Table 55.2 style).",
 "B1": "Frequency scale, close-ended because the universe of answers is known; coded 1-6 like Never...Very often.",
 "B2": "Multiple choice: more than one answer can be true (slide: types of close-ended questions).",
 "B3": "Single choice: only one 'most used' device is possible per respondent.",
 "B4": "Behaviour, not opinion: how long a visit really lasts (surveys collect habits).",
 "B5": "Behaviour in the last 7 days (concrete recall window) instead of 'which features are important to you' (leading).",
 "B6": "Grand-tour style recall of the most recent deadline; reveals workarounds such as WhatsApp groups. Includes 'missed it'.",
 "B7": "Where notifications are actually noticed, not where they are sent.",
 "B8": "Count of failures in a fixed window, coded 0-4.",
 "B9": "Workaround tools tell us what the current system fails to do.",
 "C1": "7-point Likert (Strongly disagree = 1 ... Strongly agree = 7).",
 "C2": "Inverted duplicate of C1 (slide 'Duplicate pair with inverted question'): reverse-code as 8 - C2; the pair's Cronbach alpha is computed on 1-Patterns.",
 "C3": "Likert; pairs with C4 to separate useful notifications from noise.",
 "C4": "Likert; worded so that agreement means a problem (avoids acquiescence bias in one direction).",
 "C5": "Likert; mobile vs. desktop parity.", "C6": "Likert; consistency across courses.", "C7": "Likert; trust in data freshness.",
 "C8": "Likert; self-efficacy / need for help.", "C9": "Likert; perceived accessibility.", "C10": "Likert; collaboration support.",
 "C11": "Single object rated on one dimension: 7-point satisfaction (slide: close-ended questions).",
 "C12": "Difficulty of the LAST concrete task, 5-point Very difficult...Very easy (Table 55.2).",
 "D1": "Ranking: respondents must prioritise as in real life (slide: ranking questions).",
 "E1": "Open-ended grand tour, non-leading: 'the last time it frustrated you', not 'what do you dislike'.",
 "E2": "Balances E1 (assuming only frustration would be biased framing).",
 "E3": "Wrap-up question from the interview structure slide.",
}
r = 5
for s in Q["sections"]:
    for q in s["questions"]:
        opts = ""
        if q["type"] in ("single", "multi", "rank"): opts = "\n".join(f"{o['label']}  [{o['k']}]" for o in q["options"])
        elif q["type"] == "likert": opts = "\n".join(f"{i+1} = {l}" for i, l in enumerate(Q["likert7"]))
        elif q["type"] == "scale": opts = f"{q['min']} = {q['anchors'][0]} ... {q['max']} = {q['anchors'][1]}"
        elif q["type"] == "text": opts = "Free text"
        coding = ""
        if q["type"] == "single" and any("v" in o for o in q["options"]): coding = ", ".join(f"{o['k']} = {o['v']}" for o in q["options"])
        elif q["type"] == "single": coding = "categorical"
        elif q["type"] == "multi": coding = "1 = selected, 0 = not"
        elif q["type"] in ("likert", "scale"): coding = f"{q.get('min',1)}-{q.get('max',7)}" + (" (reverse: 8 - x)" if q.get("reverse_of") else "")
        elif q["type"] == "rank": coding = "rank 1-6 per activity"
        col = f"{q['id']}_{q['key']}"
        if q["type"] in ("multi", "rank"): col += "__<option>"
        typ = {"single": "Single choice", "multi": "Multiple choice", "likert": "Likert 7-pt", "scale": f"Scale {q.get('min',1)}-{q.get('max',7)}", "rank": "Ranking", "text": "Open-ended"}[q["type"]]
        for c, v in enumerate([q["id"], s["title"], q["text"], typ, opts, coding, col, RATIONALE.get(q["id"], "")], start=1):
            cellw(ws, r, c, v, bold=(c == 1))
        ws.row_dimensions[r].height = max(30, 13 * (opts.count("\n") + 1))
        r += 1
note(ws, f"A{r+1}", "Every question can be skipped in user.html (voluntary participation); an empty cell in Responses means the respondent skipped it. Section D asks for a full ranking.", False)
ws.freeze_panes = "A5"

# =================================================================== Codes
wc = wb.create_sheet("Codes")
header_row(wc, 1, 1, ["lookup", "question", "code", "label", "value"])
r = 2
for s in Q["sections"]:
    for q in s["questions"]:
        if q["type"] in ("single", "multi", "rank"):
            for o in q["options"]:
                wc.cell(row=r, column=1, value=f"{q['id']}|{o['k']}"); wc.cell(row=r, column=2, value=q["id"]); wc.cell(row=r, column=3, value=o["k"])
                wc.cell(row=r, column=4, value=o["label"]); wc.cell(row=r, column=5, value=o.get("v"))
                for c in range(1, 6): wc.cell(row=r, column=c).font = f(9)
                r += 1
CODES_LAST = r - 1
for k, v in {"A": 18, "B": 10, "C": 14, "D": 52, "E": 8}.items(): wc.column_dimensions[k].width = v
wc.freeze_panes = "A2"

# =================================================================== Responses
wr = wb.create_sheet("Responses")
cols = list(rows[0].keys())          # survey columns incl. Segment, PC1, PC2 at the end
survey_cols = [c for c in cols if c not in ("Segment", "PC1", "PC2")]
derived = ["freq_v", "session_v", "hours_v", "lost_v", "courses_v", "phone_primary", "C2_rev"]
all_cols = survey_cols + derived + ["Segment", "PC1", "PC2"]
COL = {name: i + 1 for i, name in enumerate(all_cols)}
def R(name, absolute=True):
    l = L(COL[name]); return f"Responses!${l}$2:${l}${LAST}" if absolute else f"Responses!{l}2:{l}{LAST}"
for name, i in COL.items():
    c = wr.cell(row=1, column=i, value=name); c.font = f(9, True, "FFFFFF"); c.alignment = CENTER; c.border = BORDER
    c.fill = fill("6B7280" if name in derived else ("7C3AED" if name in ("Segment", "PC1", "PC2") else TEAL))
wr.row_dimensions[1].height = 30
numeric_prefixes = ("A9_", "C", "D1_", "A8_", "B2_", "B5_", "B6_", "B9_")
for i, row in enumerate(rows, start=2):
    for name in survey_cols:
        v = row[name]
        if name.startswith(numeric_prefixes) and v not in ("", None):
            try: v = int(v)
            except ValueError: pass
        wr.cell(row=i, column=COL[name], value=v).font = f(9)
    wr.cell(row=i, column=COL["Segment"], value=row["Segment"]).font = f(9)
    wr.cell(row=i, column=COL["PC1"], value=float(row["PC1"])).font = f(9)
    wr.cell(row=i, column=COL["PC2"], value=float(row["PC2"])).font = f(9)
# derived formula columns (rows 2..LAST) so appended real data is covered
def lookup_formula(qid, src_name, r):
    return f'=IF(${L(COL[src_name])}{r}="","",INDEX(Codes!$E$2:$E${CODES_LAST},MATCH("{qid}|"&${L(COL[src_name])}{r},Codes!$A$2:$A${CODES_LAST},0)))'
for r in range(2, LAST + 1):
    wr.cell(row=r, column=COL["freq_v"], value=lookup_formula("B1", "B1_freq", r))
    wr.cell(row=r, column=COL["session_v"], value=lookup_formula("B4", "B4_session", r))
    wr.cell(row=r, column=COL["hours_v"], value=lookup_formula("A7", "A7_hours", r))
    wr.cell(row=r, column=COL["lost_v"], value=lookup_formula("B8", "B8_lost", r))
    wr.cell(row=r, column=COL["courses_v"], value=lookup_formula("A6", "A6_courses", r))
    wr.cell(row=r, column=COL["phone_primary"], value=f'=IF(${L(COL["B3_primary"])}{r}="","",IF(${L(COL["B3_primary"])}{r}="phone",1,0))')
    wr.cell(row=r, column=COL["C2_rev"], value=f'=IF(${L(COL["C2_find_slow"])}{r}="","",8-${L(COL["C2_find_slow"])}{r})')
    for name in derived: wr.cell(row=r, column=COL[name]).font = f(9, color="374151")
wr.cell(row=1, column=COL["freq_v"]).comment = Comment("Derived columns (grey): numeric recodes via INDEX/MATCH on sheet Codes. Purple columns: k-means cluster label and PCA coordinates pasted from the analysis (Excel cannot run k-means).", "Design team")
for name, i in COL.items():
    w = 9
    if name in ("id", "submitted"): w = 11 if name == "id" else 17
    if name.startswith("E"): w = 60
    if name == "Segment": w = 24
    if name.startswith(("A1", "A2", "A4", "A5", "B1", "B3", "B4", "B7")): w = 10
    wr.column_dimensions[L(i)].width = w
wr.freeze_panes = "C2"
wr.auto_filter.ref = f"A1:{L(len(all_cols))}{N+1}"
NCELL = "'1-Patterns'!$C$4"   # N lives on the patterns sheet
SEGRNG = R("Segment")

# =================================================================== 1-Patterns
wp = wb.create_sheet("1-Patterns")
for k, v in {"A": 4, "B": 46, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 9, "J": 9, "K": 9, "L": 9, "M": 9, "N": 9, "O": 9}.items(): wp.column_dimensions[k].width = v
title(wp, "B1", "Step 2. Analyse data to find patterns", 14)
note(wp, "B2", "Every number on this sheet is a live formula over the Responses sheet (rows 2-500). Percentages are of N unless stated. Charts are on the right.", False); wp.merge_cells("B2:H2")
wp["B4"] = "N (respondents)"; wp["B4"].font = f(10, True); wp["C4"] = f"=COUNTA(Responses!$A$2:$A${LAST})"; wp["C4"].font = f(10, True, TEAL_D)
charts = []   # (anchor_row, chart)
r = 6
def dist_single(ws, r, heading, qid, colname, with_code=True):
    section_title(ws, r, heading, 8); r += 1
    header_row(ws, r, 2, ["Answer", "Code", "Count", "% of N"], fill_hex=GREY_D, color="222222"); r += 1
    r0 = r
    for o in QI[qid]["options"]:
        cellw(ws, r, 2, o["label"]); cellw(ws, r, 3, o["k"], align=CENTER)
        cellw(ws, r, 4, f'=COUNTIF({R(colname)},"{o["k"]}")', align=CENTER); cellw(ws, r, 5, f"=D{r}/{NCELL}", fmt="0%", align=CENTER); r += 1
    return r, r0, r - 1
def dist_multi(ws, r, heading, qid, prefix):
    section_title(ws, r, heading, 8); r += 1
    header_row(ws, r, 2, ["Option (select all that apply)", "Code", "Count", "% of N"], fill_hex=GREY_D, color="222222"); r += 1
    r0 = r
    for o in QI[qid]["options"]:
        cellw(ws, r, 2, o["label"]); cellw(ws, r, 3, o["k"], align=CENTER)
        cellw(ws, r, 4, f'=SUM({R(prefix + "__" + o["k"])})', align=CENTER); cellw(ws, r, 5, f"=D{r}/{NCELL}", fmt="0%", align=CENTER); r += 1
    return r, r0, r - 1
def bar(ws, title_, cats_ref, vals_ref, anchor, w=14, h=7.5, color="1B7F5E", ymax=None, ypct=True):
    ch = BarChart(); ch.type = "bar"; ch.style = 10; ch.title = title_; ch.legend = None
    ch.add_data(vals_ref, titles_from_data=False); ch.set_categories(cats_ref)
    ch.y_axis.number_format = "0%" if ypct else "0.0"; ch.y_axis.majorGridlines = None
    if ymax: ch.y_axis.scaling.max = ymax
    ch.y_axis.scaling.min = 0
    ch.width, ch.height = w, h; ch.series[0].graphicalProperties.solidFill = color; ch.series[0].graphicalProperties.line.solidFill = color
    ch.x_axis.delete = False; ch.y_axis.delete = False
    ws.add_chart(ch, anchor)

# --- A. Sample
r, a0, a1 = dist_single(wp, r, "A. Who answered", "A1", "A1_role"); r += 1
r, b0, b1 = dist_single(wp, r, "Age", "A2", "A2_age"); r += 1
r, _, _ = dist_single(wp, r, "Gender", "A3", "A3_gender"); r += 1
r, _, _ = dist_single(wp, r, "School", "A4", "A4_school"); r += 1
r, _, _ = dist_single(wp, r, "Secondary education", "A5", "A5_origin"); r += 1
r, h0, h1 = dist_single(wp, r, "Hours per week on commitments outside study/job (A7)", "A7", "A7_hours"); r += 1
r, x0, x1 = dist_multi(wp, r, "Accessibility needs (A8)", "A8", "A8_access"); r += 1
section_title(wp, r, "Comfort learning new software on your own (A9, 1-7)", 8); r += 1
header_row(wp, r, 2, ["Statistic", "Value"], fill_hex=GREY_D, color="222222"); r += 1
for lab, fm in [("Mean", f"=AVERAGE({R('A9_tech')})"), ("SD", f"=STDEV({R('A9_tech')})"), ("Median", f"=MEDIAN({R('A9_tech')})"), ("% scoring 5-7 (comfortable)", f'=COUNTIF({R("A9_tech")},">=5")/{NCELL}'), ("% scoring 1-3 (uncomfortable)", f'=COUNTIF({R("A9_tech")},"<=3")/{NCELL}')]:
    cellw(wp, r, 2, lab); cellw(wp, r, 3, fm, fmt="0%" if "%" in lab else "0.00", align=CENTER); r += 1
r += 1
# --- B. Habits
r, f0, f1 = dist_single(wp, r, "B. How often do people open Canvas (B1)", "B1", "B1_freq")
cellw(wp, r, 2, "Mean frequency score (1-6)", True); cellw(wp, r, 3, f"=AVERAGE({R('freq_v')})", fmt="0.00", align=CENTER); r += 1
cellw(wp, r, 2, "% at least once a day", True); cellw(wp, r, 3, f'=(COUNTIF({R("B1_freq")},"daily")+COUNTIF({R("B1_freq")},"several"))/{NCELL}', fmt="0%", align=CENTER); r += 2
bar(wp, "How often do respondents open Canvas?", Reference(wp, min_col=2, min_row=f0, max_row=f1), Reference(wp, min_col=5, min_row=f0, max_row=f1), "J6")
r, d0, d1 = dist_multi(wp, r, "Devices used in the last 7 days (B2)", "B2", "B2_devices"); r += 1
r, p0, p1 = dist_single(wp, r, "Most used device (B3)", "B3", "B3_primary"); r += 1
r, s0, s1 = dist_single(wp, r, "Typical visit length (B4)", "B4", "B4_session")
cellw(wp, r, 2, "% of visits under 5 minutes", True); cellw(wp, r, 3, f'=(COUNTIF({R("B4_session")},"lt2")+COUNTIF({R("B4_session")},"2-5"))/{NCELL}', fmt="0%", align=CENTER); r += 2
r, fe0, fe1 = dist_multi(wp, r, "Parts of Canvas used in the last 7 days (B5)", "B5", "B5_features"); r += 1
bar(wp, "Parts of Canvas used in the last 7 days", Reference(wp, min_col=2, min_row=fe0, max_row=fe1 - 1), Reference(wp, min_col=5, min_row=fe0, max_row=fe1 - 1), "J22", h=9)
r, dl0, dl1 = dist_multi(wp, r, "How people found out about their most recent deadline (B6)", "B6", "B6_deadline"); r += 1
bar(wp, "How did you first find out about your last deadline?", Reference(wp, min_col=2, min_row=dl0, max_row=dl1), Reference(wp, min_col=5, min_row=dl0, max_row=dl1), "J41", h=8, color="B45309")
r, _, _ = dist_single(wp, r, "Where notifications are noticed (B7)", "B7", "B7_notif"); r += 1
r, _, _ = dist_single(wp, r, "Times failed to find something in the last 7 days (B8)", "B8", "B8_lost")
cellw(wp, r, 2, "% who failed at least twice", True); cellw(wp, r, 3, f'=(COUNTIF({R("B8_lost")},"2-3")+COUNTIF({R("B8_lost")},"4-6")+COUNTIF({R("B8_lost")},"7+"))/{NCELL}', fmt="0%", align=CENTER); r += 2
r, _, _ = dist_multi(wp, r, "Other tools used to keep track of courses (B9)", "B9", "B9_tools"); r += 1

# --- C. Attitudes
section_title(wp, r, "C. Experience statements (1 = strongly disagree ... 7 = strongly agree)", 15); r += 1
header_row(wp, r, 2, ["Statement", "Item", "Mean", "SD", "Median", "% disagree (1-3)", "% neutral (4)", "% agree (5-7)", "n=1", "n=2", "n=3", "n=4", "n=5", "n=6", "n=7"], fill_hex=GREY_D, color="222222"); r += 1
lk0 = r
LIK = ["C1_find_fast", "C2_find_slow", "C3_notif_useful", "C4_notif_noise", "C5_mobile", "C6_consistent", "C7_uptodate", "C8_confident", "C9_accessible", "C10_group"]
for k in LIK:
    rng = R(k)
    cellw(wp, r, 2, A["likert_text"][k]); cellw(wp, r, 3, k.split("_")[0], align=CENTER)
    cellw(wp, r, 4, f"=AVERAGE({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 5, f"=STDEV({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 6, f"=MEDIAN({rng})", fmt="0.0", align=CENTER)
    cellw(wp, r, 7, f'=COUNTIF({rng},"<=3")/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 8, f'=COUNTIF({rng},4)/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 9, f'=COUNTIF({rng},">=5")/{NCELL}', fmt="0%", align=CENTER)
    for j in range(1, 8): cellw(wp, r, 9 + j, f"=COUNTIF({rng},{j})", align=CENTER)
    r += 1
lk1 = r - 1
wp.conditional_formatting.add(f"D{lk0}:D{lk1}", ColorScaleRule(start_type="num", start_value=1, start_color="F8B4B4", mid_type="num", mid_value=4, mid_color="FFFFFF", end_type="num", end_value=7, end_color="9FD8B8"))
bar(wp, "Mean agreement per statement (1-7)", Reference(wp, min_col=3, min_row=lk0, max_row=lk1), Reference(wp, min_col=4, min_row=lk0, max_row=lk1), "J60", h=8, ypct=False, ymax=7, color="4C1D95")
r += 1
header_row(wp, r, 2, ["Overall ratings", "Item", "Mean", "SD", "Median", "% low", "% mid", "% high"], fill_hex=GREY_D, color="222222"); r += 1
cellw(wp, r, 2, "Overall satisfaction (1 = extremely dissatisfied ... 7 = extremely satisfied)"); cellw(wp, r, 3, "C11", align=CENTER)
rng = R("C11_satisfaction"); cellw(wp, r, 4, f"=AVERAGE({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 5, f"=STDEV({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 6, f"=MEDIAN({rng})", fmt="0.0", align=CENTER)
cellw(wp, r, 7, f'=COUNTIF({rng},"<=3")/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 8, f'=COUNTIF({rng},4)/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 9, f'=COUNTIF({rng},">=5")/{NCELL}', fmt="0%", align=CENTER); r += 1
cellw(wp, r, 2, "Difficulty of the last task (1 = very difficult ... 5 = very easy)"); cellw(wp, r, 3, "C12", align=CENTER)
rng = R("C12_ease"); cellw(wp, r, 4, f"=AVERAGE({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 5, f"=STDEV({rng})", fmt="0.00", align=CENTER); cellw(wp, r, 6, f"=MEDIAN({rng})", fmt="0.0", align=CENTER)
cellw(wp, r, 7, f'=COUNTIF({rng},"<=2")/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 8, f'=COUNTIF({rng},3)/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 9, f'=COUNTIF({rng},">=4")/{NCELL}', fmt="0%", align=CENTER); r += 2

section_title(wp, r, "Reliability check of the inverted pair C1 / C2 (slide: duplicate pair with inverted question)", 8); r += 1
cellw(wp, r, 2, "Pearson r between C1 and reversed C2 (8 - C2)"); cellw(wp, r, 3, f"=CORREL({R('C1_find_fast')},{R('C2_rev')})", fmt="0.00", align=CENTER); r += 1
cellw(wp, r, 2, "Cronbach's alpha for the 2-item scale  = 2r / (1 + r)"); cellw(wp, r, 3, f"=2*C{r-1}/(1+C{r-1})", fmt="0.00", align=CENTER); r += 1
cellw(wp, r, 2, "Mean 'findability' score = average of C1 and reversed C2"); cellw(wp, r, 3, f"=(AVERAGE({R('C1_find_fast')})+AVERAGE({R('C2_rev')}))/2", fmt="0.00", align=CENTER); r += 1
note(wp, f"B{r}", "Alpha above 0.7 means respondents answered the pair consistently; the two items can be averaged into one 'findability' score. A respondent who agrees with both C1 and C2 is a candidate straight-liner."); wp.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8); wp.row_dimensions[r].height = 28; r += 2

# --- D. Ranking
section_title(wp, r, "D. Priorities: rank of six activities (1 = most important)", 8); r += 1
header_row(wp, r, 2, ["Activity", "Code", "Mean rank", "% ranked 1st", "% ranked top-2", "% ranked last"], fill_hex=GREY_D, color="222222"); r += 1
rk0 = r
for o in QI["D1"]["options"]:
    rng = R("D1_rank__" + o["k"])
    cellw(wp, r, 2, o["label"]); cellw(wp, r, 3, o["k"], align=CENTER); cellw(wp, r, 4, f"=AVERAGE({rng})", fmt="0.00", align=CENTER)
    cellw(wp, r, 5, f"=COUNTIF({rng},1)/{NCELL}", fmt="0%", align=CENTER); cellw(wp, r, 6, f'=COUNTIF({rng},"<=2")/{NCELL}', fmt="0%", align=CENTER); cellw(wp, r, 7, f"=COUNTIF({rng},6)/{NCELL}", fmt="0%", align=CENTER); r += 1
rk1 = r - 1
wp.conditional_formatting.add(f"D{rk0}:D{rk1}", ColorScaleRule(start_type="min", start_color="9FD8B8", end_type="max", end_color="FFFFFF"))
r += 1

# --- E. Cross-tabs
section_title(wp, r, "E. Cross-tabulations: do different users behave differently?", 8); r += 1
roles = QI["A1"]["options"]; freqs = QI["B1"]["options"]
header_row(wp, r, 2, ["Role x frequency (counts)"] + [o["label"] for o in freqs] + ["Total"], fill_hex=GREY_D, color="222222"); r += 1
for ro in roles:
    cellw(wp, r, 2, ro["label"])
    for j, fo in enumerate(freqs):
        cellw(wp, r, 3 + j, f'=COUNTIFS({R("A1_role")},"{ro["k"]}",{R("B1_freq")},"{fo["k"]}")', align=CENTER)
    cellw(wp, r, 3 + len(freqs), f"=SUM(C{r}:{L(2+len(freqs))}{r})", True, align=CENTER); r += 1
r += 1
header_row(wp, r, 2, ["Mean by role", "n", "Satisfaction (1-7)", "Find quickly C1", "Confident C8", "Mobile parity C5", "Notif. noise C4"], fill_hex=GREY_D, color="222222"); r += 1
for ro in roles:
    cellw(wp, r, 2, ro["label"]); cellw(wp, r, 3, f'=COUNTIF({R("A1_role")},"{ro["k"]}")', align=CENTER)
    for j, k in enumerate(["C11_satisfaction", "C1_find_fast", "C8_confident", "C5_mobile", "C4_notif_noise"]):
        cellw(wp, r, 4 + j, f'=IFERROR(AVERAGEIFS({R(k)},{R("A1_role")},"{ro["k"]}"),"")', fmt="0.00", align=CENTER)
    r += 1
r += 1
header_row(wp, r, 2, ["Mean by most-used device", "n", "Satisfaction (1-7)", "Mobile parity C5", "Visit length (1-5)", "Frequency (1-6)"], fill_hex=GREY_D, color="222222"); r += 1
for po in QI["B3"]["options"]:
    cellw(wp, r, 2, po["label"]); cellw(wp, r, 3, f'=COUNTIF({R("B3_primary")},"{po["k"]}")', align=CENTER)
    for j, k in enumerate(["C11_satisfaction", "C5_mobile", "session_v", "freq_v"]):
        cellw(wp, r, 4 + j, f'=IFERROR(AVERAGEIFS({R(k)},{R("B3_primary")},"{po["k"]}"),"")', fmt="0.00", align=CENTER)
    r += 1
r += 1
header_row(wp, r, 2, ["Mean by frequency of use", "n", "Satisfaction (1-7)", "Find quickly C1", "% found out late (B6)"], fill_hex=GREY_D, color="222222"); r += 1
for fo in freqs:
    cellw(wp, r, 2, fo["label"]); cellw(wp, r, 3, f'=COUNTIF({R("B1_freq")},"{fo["k"]}")', align=CENTER)
    cellw(wp, r, 4, f'=IFERROR(AVERAGEIFS({R("C11_satisfaction")},{R("B1_freq")},"{fo["k"]}"),"")', fmt="0.00", align=CENTER)
    cellw(wp, r, 5, f'=IFERROR(AVERAGEIFS({R("C1_find_fast")},{R("B1_freq")},"{fo["k"]}"),"")', fmt="0.00", align=CENTER)
    cellw(wp, r, 6, f'=IFERROR(COUNTIFS({R("B6_deadline__late")},1,{R("B1_freq")},"{fo["k"]}")/C{r},"")', fmt="0%", align=CENTER); r += 1
r += 1
header_row(wp, r, 2, ["Outside commitments (A7) x missed deadline", "n", "% found out late (B6)", "Mean frequency (1-6)", "Satisfaction (1-7)"], fill_hex=GREY_D, color="222222"); r += 1
for ho in QI["A7"]["options"]:
    cellw(wp, r, 2, ho["label"]); cellw(wp, r, 3, f'=COUNTIF({R("A7_hours")},"{ho["k"]}")', align=CENTER)
    cellw(wp, r, 4, f'=IFERROR(COUNTIFS({R("B6_deadline__late")},1,{R("A7_hours")},"{ho["k"]}")/C{r},"")', fmt="0%", align=CENTER)
    cellw(wp, r, 5, f'=IFERROR(AVERAGEIFS({R("freq_v")},{R("A7_hours")},"{ho["k"]}"),"")', fmt="0.00", align=CENTER)
    cellw(wp, r, 6, f'=IFERROR(AVERAGEIFS({R("C11_satisfaction")},{R("A7_hours")},"{ho["k"]}"),"")', fmt="0.00", align=CENTER); r += 1
r += 1

# --- F. Correlations
section_title(wp, r, "F. Correlations between experience items (Pearson r)", 15); r += 1
ITEMS = LIK + ["C11_satisfaction", "C12_ease"]
header_row(wp, r, 2, ["Item"] + [k.split("_")[0] for k in ITEMS], fill_hex=GREY_D, color="222222"); r += 1
co0 = r
for a in ITEMS:
    cellw(wp, r, 2, f'{a.split("_")[0]}  {A["likert_text"][a]}'[:60])
    for j, b in enumerate(ITEMS):
        cellw(wp, r, 3 + j, 1 if a == b else f"=CORREL({R(a)},{R(b)})", fmt="0.00", align=CENTER)
    r += 1
wp.conditional_formatting.add(f"C{co0}:{L(2+len(ITEMS))}{r-1}", ColorScaleRule(start_type="num", start_value=-1, start_color="F8B4B4", mid_type="num", mid_value=0, mid_color="FFFFFF", end_type="num", end_value=1, end_color="9FD8B8"))
note(wp, f"B{r}", "Read: C1 and C2 should be strongly negative (they are the inverted pair). Look for what travels with satisfaction (C11)."); r += 2

# --- G. Key patterns (narrative)
O = A["overall"]; lm = O["likert_mean"]
section_title(wp, r, "G. Key patterns we take into the segmentation (written from the numbers above)", 8); r += 1
pat = [
 f"1. Canvas is a daily habit for most: {O['pct_several_or_daily']:.0f}% open it at least once a day, but visits are short for a third of respondents ({O['pct_short_session']:.0f}% under 5 minutes). Frequency and satisfaction rise together (several-times-a-day users: {A['sat_by_freq']['Several times a day']['mean']:.1f}/7 vs once-a-week users: {A['sat_by_freq']['About once a week']['mean']:.1f}/7).",
 f"2. Findability is the weakest point: 'I can quickly find the material I need' averages {lm['C1_find_fast']:.1f}/7 and 'course pages are organised similarly' only {lm['C6_consistent']:.1f}/7; {O['pct_lost_2plus']:.0f}% failed to find something at least twice last week. Findability is the item most correlated with overall satisfaction (r = {A['corr']['C1_find_fast']['C11_satisfaction']:.2f}).",
 f"3. Notifications are noise, not signal: 'I receive notifications that are not relevant' averages {lm['C4_notif_noise']:.1f}/7 while 'notifications help me stay on top of what is due' only {lm['C3_notif_useful']:.1f}/7. {O['deadline']['A classmate or colleague told me (e.g. WhatsApp)']:.0f}% learned of their last deadline from a peer and {O['deadline']['My own calendar or planner']:.0f}% from their own calendar; only {O['deadline']['Canvas calendar or To-do list']:.0f}% from the Canvas calendar. {O['pct_missed_deadline']:.0f}% found out late or missed it.",
 f"4. Mobile is a second-class citizen: mobile parity scores {lm['C5_mobile']:.1f}/7 overall and {A['mobile_by_primary']['Smartphone']:.1f}/7 among the {O['primary']['Smartphone']} people whose main device IS a smartphone, yet those users are the most satisfied group ({A['sat_by_primary']['Smartphone']['mean']:.1f}/7) because they use it for quick checks (deadlines, grades).",
 f"5. Two priorities dominate the ranking: finding materials (mean rank {O['rank_mean']['Finding course materials']:.1f}) and keeping track of deadlines ({O['rank_mean']['Keeping track of deadlines']:.1f}); organising group work is last ({O['rank_mean']['Organising group work']:.1f}) even though 'working with others is well supported' scores only {lm['C10_group']:.1f}/7.",
 f"6. Roles differ: lecturers/professors are the least satisfied role ({A['sat_by_role']['Lecturer / professor']['mean']:.1f}/7) and split into all-day users and once-a-week users (see role x frequency); confidence ('without asking for help') ranges from very high to very low within the same role. Behaviour, not job title, will drive the segments.",
 f"7. Extreme users surface real barriers: respondents with accessibility needs rate 'easy to read, see and hear' far below the rest; respondents with 17+ hours of outside commitments open Canvas less often and miss deadlines more often (see cross-tab A7).",
 f"8. Everyone works around the system: {O['tools']['Phone or computer calendar (Google, Outlook, Apple)']:.0f}% use their own calendar, {O['tools']['WhatsApp, Signal or Discord group']:.0f}% a chat group, {O['tools']['Notion, Obsidian or OneNote']:.0f}% a notes app; nobody (0%) relies on Canvas alone.",
]
for p in pat:
    c = wp.cell(row=r, column=2, value=p); c.font = f(10); c.alignment = WRAP; wp.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9); wp.row_dimensions[r].height = 44; r += 1
note(wp, f"B{r}", "Numbers in this narrative were read from the formulas above when the sheet was written; if you replace the data, re-read them."); r += 1
wp.freeze_panes = "A5"

# =================================================================== 2-Segments
wsg = wb.create_sheet("2-Segments")
for k, v in {"A": 4, "B": 44, "C": 15, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 13, "J": 4, "K": 12, "L": 12, "M": 12}.items(): wsg.column_dimensions[k].width = v
title(wsg, "B1", "Step 3. Cluster user segments", 14)
note(wsg, "B2", "k-means on standardised features (frequency, visit length, phone as main device, tech comfort, outside commitments, times lost, the 12 experience items and the 6 ranks). Cluster labels are pasted in Responses!Segment; every profile number below is a live formula on that column.", False)
wsg.merge_cells("B2:I2"); wsg.row_dimensions[2].height = 40
r = 4
section_title(wsg, r, "Choosing k (elbow and silhouette; k = 6 chosen)", 9); r += 1
header_row(wsg, r, 2, ["k", "Inertia (within-cluster SS)", "Silhouette"], fill_hex=GREY_D, color="222222"); r += 1
k0 = r
for k, ine, sil in zip(A["k_selection"]["k"], A["k_selection"]["inertia"], A["k_selection"]["silhouette"]):
    cellw(wsg, r, 2, k, align=CENTER, fill_hex=(LAV if k == 6 else None)); cellw(wsg, r, 3, ine, fmt="0.0", align=CENTER); cellw(wsg, r, 4, sil, fmt="0.000", align=CENTER); r += 1
k1 = r - 1
note(wsg, f"B{r}", "The elbow flattens after k = 5-6 and the silhouette is flat between 5 and 6. We chose 6 because the sixth cluster isolates the accessibility extreme users, who are few but design-critical; with k = 5 they were absorbed into the Structured Scholars. (Values pasted from the Python/vis.html run.)")
wsg.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9); wsg.row_dimensions[r].height = 40; r += 1
lc = LineChart(); lc.title = "Inertia by k"; lc.style = 12; lc.legend = None; lc.add_data(Reference(wsg, min_col=3, min_row=k0, max_row=k1), titles_from_data=False); lc.set_categories(Reference(wsg, min_col=2, min_row=k0, max_row=k1)); lc.width, lc.height = 9, 6; lc.x_axis.title = "k"; lc.x_axis.delete = False; lc.y_axis.delete = False
wsg.add_chart(lc, "K4")
lc2 = LineChart(); lc2.title = "Silhouette by k"; lc2.style = 12; lc2.legend = None; lc2.add_data(Reference(wsg, min_col=4, min_row=k0, max_row=k1), titles_from_data=False); lc2.set_categories(Reference(wsg, min_col=2, min_row=k0, max_row=k1)); lc2.width, lc2.height = 9, 6; lc2.x_axis.title = "k"; lc2.x_axis.delete = False; lc2.y_axis.delete = False
wsg.add_chart(lc2, "K17")
r += 1
section_title(wsg, r, "Segment sizes", 9); r += 1
header_row(wsg, r, 2, ["Segment", "n", "% of N", "One-line description"], fill_hex=GREY_D, color="222222"); r += 1
sz0 = r
SEGCELL = {}
for s in SEGS:
    cellw(wsg, r, 2, s, True); cellw(wsg, r, 3, f'=COUNTIF({SEGRNG},"{s}")', align=CENTER); cellw(wsg, r, 4, f"=C{r}/{NCELL}", fmt="0%", align=CENTER)
    cellw(wsg, r, 5, A["segments"][s]["tagline"]); wsg.merge_cells(start_row=r, start_column=5, end_row=r, end_column=9); wsg.row_dimensions[r].height = 30
    SEGCELL[s] = f"'2-Segments'!$C${r}"; r += 1
r += 1
# ---- profile matrix
section_title(wsg, r, "Segment profile matrix (live formulas; colour scale compares segments within each row)", 9); r += 1
header_row(wsg, r, 2, ["Feature"] + SEGS + ["All"], fill_hex=TEAL); r += 1
def seg_avg(col, s): return f'=IFERROR(AVERAGEIFS({R(col)},{SEGRNG},"{s}"),"")'
def seg_pct_eq(col, val, s): return f'=IFERROR(COUNTIFS({R(col)},{val},{SEGRNG},"{s}")/{SEGCELL[s]},"")'
def all_avg(col): return f"=AVERAGE({R(col)})"
def all_pct_eq(col, val): return f"=COUNTIF({R(col)},{val})/{NCELL}"
def prow(label, fn_seg, fn_all, fmt):
    global r
    cellw(wsg, r, 2, label)
    for j, s in enumerate(SEGS): cellw(wsg, r, 3 + j, fn_seg(s), fmt=fmt, align=CENTER)
    cellw(wsg, r, 3 + len(SEGS), fn_all(), fmt=fmt, align=CENTER, fill_hex=GREY)
    wsg.conditional_formatting.add(f"C{r}:{L(2+len(SEGS))}{r}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="7FCBA4"))
    r += 1
def subhead(label):
    global r
    c = wsg.cell(row=r, column=2, value=label); c.font = f(10, True, TEAL_D)
    for cc in range(2, 3 + len(SEGS) + 1): wsg.cell(row=r, column=cc).fill = fill(GREY)
    r += 1
subhead("Behaviour")
prow("Opens Canvas (1 = <1x/week ... 6 = several times a day)", lambda s: seg_avg("freq_v", s), lambda: all_avg("freq_v"), "0.00")
prow("% opening at least once a day", lambda s: f'=IFERROR((COUNTIFS({R("B1_freq")},"daily",{SEGRNG},"{s}")+COUNTIFS({R("B1_freq")},"several",{SEGRNG},"{s}"))/{SEGCELL[s]},"")', lambda: f'=(COUNTIF({R("B1_freq")},"daily")+COUNTIF({R("B1_freq")},"several"))/{NCELL}', "0%")
prow("Visit length (1 = <2 min ... 5 = >30 min)", lambda s: seg_avg("session_v", s), lambda: all_avg("session_v"), "0.00")
prow("% smartphone is the most used device", lambda s: seg_pct_eq("B3_primary", '"phone"', s), lambda: all_pct_eq("B3_primary", '"phone"'), "0%")
prow("% used the Canvas app last week", lambda s: seg_pct_eq("B2_devices__phone_app", 1, s), lambda: all_pct_eq("B2_devices__phone_app", 1), "0%")
prow("Comfort learning software alone (1-7)", lambda s: seg_avg("A9_tech", s), lambda: all_avg("A9_tech"), "0.00")
prow("Outside commitments (0 = none ... 4 = >24 h/week)", lambda s: seg_avg("hours_v", s), lambda: all_avg("hours_v"), "0.00")
prow("% with 17+ h/week outside commitments", lambda s: f'=IFERROR((COUNTIFS({R("A7_hours")},"17-24",{SEGRNG},"{s}")+COUNTIFS({R("A7_hours")},"25+",{SEGRNG},"{s}"))/{SEGCELL[s]},"")', lambda: f'=(COUNTIF({R("A7_hours")},"17-24")+COUNTIF({R("A7_hours")},"25+"))/{NCELL}', "0%")
ACC_KEYS = ["screenreader", "captions", "dyslexia", "motor", "contrast"]
def acc_seg(s):   # respondents who ticked neither 'none' nor 'prefer not to say' have at least one need
    return '=IFERROR(COUNTIFS(%s,0,%s,0,%s,"%s")/%s,"")' % (R("A8_access__none"), R("A8_access__na"), SEGRNG, s, SEGCELL[s])
def acc_all():
    return "=COUNTIFS(%s,0,%s,0)/%s" % (R("A8_access__none"), R("A8_access__na"), NCELL)
prow("% with an accessibility need (A8, not 'none'/'prefer not')", acc_seg, acc_all, "0%")
prow("Times failed to find something last week (0-4)", lambda s: seg_avg("lost_v", s), lambda: all_avg("lost_v"), "0.00")
prow("% found out about last deadline late / missed it", lambda s: seg_pct_eq("B6_deadline__late", 1, s), lambda: all_pct_eq("B6_deadline__late", 1), "0%")
prow("Courses this block (1-6)", lambda s: seg_avg("courses_v", s), lambda: all_avg("courses_v"), "0.00")
subhead("Parts of Canvas used last week (% of segment)")
for o in QI["B5"]["options"][:-1]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq(f"B5_features__{k}", 1, s), lambda k=o["k"]: all_pct_eq(f"B5_features__{k}", 1), "0%")
subhead("How they found out about the last deadline (% of segment)")
for o in QI["B6"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq(f"B6_deadline__{k}", 1, s), lambda k=o["k"]: all_pct_eq(f"B6_deadline__{k}", 1), "0%")
subhead("Where notifications are noticed (% of segment)")
for o in QI["B7"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq("B7_notif", f'"{k}"', s), lambda k=o["k"]: all_pct_eq("B7_notif", f'"{k}"'), "0%")
subhead("Other tools used (% of segment)")
for o in QI["B9"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq(f"B9_tools__{k}", 1, s), lambda k=o["k"]: all_pct_eq(f"B9_tools__{k}", 1), "0%")
subhead("Experience (mean, 1-7; C12 is 1-5)")
for k in ITEMS:
    prow(f'{k.split("_")[0]}  {A["likert_text"][k]}', lambda s, k=k: seg_avg(k, s), lambda k=k: all_avg(k), "0.00")
subhead("Priorities (mean rank, 1 = most important)")
for o in QI["D1"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_avg(f"D1_rank__{k}", s), lambda k=o["k"]: all_avg(f"D1_rank__{k}"), "0.00")
subhead("Who is in the segment (% of segment by role)")
for o in QI["A1"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq("A1_role", f'"{k}"', s), lambda k=o["k"]: all_pct_eq("A1_role", f'"{k}"'), "0%")
subhead("Age (% of segment)")
for o in QI["A2"]["options"]:
    prow(o["label"], lambda s, k=o["k"]: seg_pct_eq("A2_age", f'"{k}"', s), lambda k=o["k"]: all_pct_eq("A2_age", f'"{k}"'), "0%")
r += 1
# ---- role x segment
section_title(wsg, r, "Roles cut across segments (counts): personas follow behaviour patterns, not job titles (About Face)", 9); r += 1
header_row(wsg, r, 2, ["Role"] + SEGS + ["Total"], fill_hex=GREY_D, color="222222"); r += 1
for o in QI["A1"]["options"]:
    cellw(wsg, r, 2, o["label"])
    for j, s in enumerate(SEGS): cellw(wsg, r, 3 + j, f'=COUNTIFS({R("A1_role")},"{o["k"]}",{SEGRNG},"{s}")', align=CENTER)
    cellw(wsg, r, 3 + len(SEGS), f"=SUM(C{r}:{L(2+len(SEGS))}{r})", True, align=CENTER); r += 1
r += 1
# ---- descriptions
section_title(wsg, r, "Segment descriptions (who, defining behaviour, key pains with evidence, a voice from the data, design implication)", 9); r += 1
header_row(wsg, r, 2, ["Segment", "Who they are", "Defining behaviour", "Key pains (evidence)", "In their own words (E1)", "Design implication for iCanvas"], fill_hex=GREY_D, color="222222"); r += 1
S = A["segments"]
def top_roles(s, n=2):
    it = sorted(S[s]["roles"].items(), key=lambda x: -x[1])[:n]; return ", ".join(f"{k.split(' (')[0]} ({v})" for k, v in it)
DESC = {
 "S1 Deadline Sprinters": (
   f"{top_roles('S1 Deadline Sprinters')}. Mostly 17-24, {S['S1 Deadline Sprinters']['pct_work17plus']:.0f}% work 17+ h. n = {S['S1 Deadline Sprinters']['n']}.",
   f"Open Canvas several times a day ({S['S1 Deadline Sprinters']['pct_several_or_daily']:.0f}% at least daily), on the phone ({S['S1 Deadline Sprinters']['pct_phone_primary']:.0f}% smartphone as main device, {S['S1 Deadline Sprinters']['devices']['Smartphone, in the Canvas app']:.0f}% use the app), for under 5 minutes ({S['S1 Deadline Sprinters']['pct_short_session']:.0f}%). Check grades ({S['S1 Deadline Sprinters']['features']['Grades / feedback']:.0f}%) and announcements. Rank deadlines #1 (mean {S['S1 Deadline Sprinters']['rank_mean']['Keeping track of deadlines']:.1f}) and grades #2.",
   f"Mobile parity {S['S1 Deadline Sprinters']['likert_mean']['C5_mobile']:.1f}/7 (lowest of any student segment); notification noise {S['S1 Deadline Sprinters']['likert_mean']['C4_notif_noise']:.1f}/7; {S['S1 Deadline Sprinters']['deadline']['A classmate or colleague told me (e.g. WhatsApp)']:.0f}% heard of their last deadline from a peer, {S['S1 Deadline Sprinters']['tools']['WhatsApp, Signal or Discord group']:.0f}% run a course chat group; trust in up-to-date grades only {S['S1 Deadline Sprinters']['likert_mean']['C7_uptodate']:.1f}/7.",
   "I wanted to quickly check if the lecture was cancelled. Had to go dashboard > course > announcements > scroll. Three taps too many when you are on the bus.",
   "A phone-first 'what changed / what is due' surface; one trustworthy push per real change; grades with context."),
 "S2 Structured Scholars": (
   f"{top_roles('S2 Structured Scholars')}. 21-29, highest tech comfort ({S['S2 Structured Scholars']['numeric_mean']['A9_tech']:.1f}/7). n = {S['S2 Structured Scholars']['n']}.",
   f"Laptop only ({S['S2 Structured Scholars']['pct_phone_primary']:.0f}% phone), about once a day, long visits ({S['S2 Structured Scholars']['session']['More than 30 minutes']} of {S['S2 Structured Scholars']['n']} over 30 min). Live in Modules ({S['S2 Structured Scholars']['features']['Modules or Files (course materials)']:.0f}%) and recordings ({S['S2 Structured Scholars']['features']['Lecture recordings (video)']:.0f}%). Rank materials #1 ({S['S2 Structured Scholars']['rank_mean']['Finding course materials']:.1f}) and recordings #2. {S['S2 Structured Scholars']['tools']['Notion, Obsidian or OneNote']:.0f}% keep a Notion/Obsidian/OneNote system.",
   f"Consistency across courses {S['S2 Structured Scholars']['likert_mean']['C6_consistent']:.1f}/7 and findability {S['S2 Structured Scholars']['likert_mean']['C1_find_fast']:.1f}/7 (both far below average); {S['S2 Structured Scholars']['pct_lost_2plus']:.0f}% failed to find something 2+ times last week; only {S['S2 Structured Scholars']['deadline']['Canvas calendar or To-do list']:.0f}% learned of the last deadline from the Canvas calendar because undated items never appear.",
   "I was looking for the reading list for week 4. One lecturer puts readings in Modules, another in Files, another in a PDF under Syllabus. I spent 15 minutes clicking around.",
   "Enforced course structure, cross-course search (incl. inside PDFs and transcripts), reliable calendar feed."),
 "S3 Overloaded Jugglers": (
   f"{top_roles('S3 Overloaded Jugglers')}. Older students (25-49), {S['S3 Overloaded Jugglers']['pct_work17plus']:.0f}% have 17+ h/week of work or care. n = {S['S3 Overloaded Jugglers']['n']}.",
   f"Open Canvas only 2-3 times a week (mean {S['S3 Overloaded Jugglers']['numeric_mean']['freq_v']:.1f} on the 1-6 scale; {S['S3 Overloaded Jugglers']['pct_several_or_daily']:.0f}% daily), in batches of 5-15 minutes; {S['S3 Overloaded Jugglers']['notif']['In my e-mail inbox']} of {S['S3 Overloaded Jugglers']['n']} rely on e-mail; recordings are essential ({S['S3 Overloaded Jugglers']['features']['Lecture recordings (video)']:.0f}%). Rank deadlines #1 and recordings #2.",
   f"{S['S3 Overloaded Jugglers']['pct_missed_deadline']:.0f}% found out about their last deadline late or missed it (vs {A['overall']['pct_missed_deadline']:.0f}% overall); notification noise {S['S3 Overloaded Jugglers']['likert_mean']['C4_notif_noise']:.1f}/7 while usefulness only {S['S3 Overloaded Jugglers']['likert_mean']['C3_notif_useful']:.1f}/7; satisfaction {S['S3 Overloaded Jugglers']['likert_mean']['C11_satisfaction']:.1f}/7 ({S['S3 Overloaded Jugglers']['pct_dissatisfied']:.0f}% dissatisfied).",
   "I can only study on Sunday. I opened Canvas and had 13 unread announcements from 4 courses. No way to see which ones actually require me to do something.",
   "A 'what did I miss / what must I do before the tutorial' digest; deadline changes that reach e-mail as one clear message."),
 "S4 Power Instructors": (
   f"{top_roles('S4 Power Instructors', 3)}. 25-49. n = {S['S4 Power Instructors']['n']}.",
   f"In Canvas all day ({S['S4 Power Instructors']['pct_several_or_daily']:.0f}% at least daily, {S['S4 Power Instructors']['session']['More than 30 minutes']} of {S['S4 Power Instructors']['n']} visits over 30 min) on desktop/laptop; use assignments, gradebook, announcements, modules and Inbox ({S['S4 Power Instructors']['features']['Inbox (messages)']:.0f}%). Rank grading #1 ({S['S4 Power Instructors']['rank_mean']['Seeing grades and feedback']:.1f}). Highest confidence ({S['S4 Power Instructors']['likert_mean']['C8_confident']:.1f}/7).",
   f"Notification noise {S['S4 Power Instructors']['likert_mean']['C4_notif_noise']:.1f}/7 (one e-mail per submission); group work support {S['S4 Power Instructors']['likert_mean']['C10_group']:.1f}/7; mobile {S['S4 Power Instructors']['likert_mean']['C5_mobile']:.1f}/7; last task difficulty {S['S4 Power Instructors']['likert_mean']['C12_ease']:.1f}/5 despite expertise: repetitive clicks (course copy dates, SpeedGrader reloads, section overrides).",
   "Grading 180 submissions in SpeedGrader: the page reloads for every student and the rubric panel collapses each time. Two clicks that cost me hours.",
   "Bulk actions (dates, messages, rubrics), groups that respect sections, per-course notification rules, analytics per tutorial group."),
 "S5 Reluctant Instructors": (
   f"{top_roles('S5 Reluctant Instructors')}. 40+, lowest tech comfort ({S['S5 Reluctant Instructors']['numeric_mean']['A9_tech']:.1f}/7). n = {S['S5 Reluctant Instructors']['n']}.",
   f"Open Canvas about once a week (mean {S['S5 Reluctant Instructors']['numeric_mean']['freq_v']:.1f}/6; {S['S5 Reluctant Instructors']['pct_several_or_daily']:.0f}% daily) on a desktop, mainly to post announcements and files; {S['S5 Reluctant Instructors']['tools']['E-mail']:.0f}% run their course through e-mail and {S['S5 Reluctant Instructors']['tools']['Paper planner or notebook']:.0f}% a paper agenda; delegate to TAs and the education office.",
   f"Confidence 'without asking for help' {S['S5 Reluctant Instructors']['likert_mean']['C8_confident']:.1f}/7 (lowest of all); last task difficulty {S['S5 Reluctant Instructors']['likert_mean']['C12_ease']:.1f}/5; satisfaction {S['S5 Reluctant Instructors']['likert_mean']['C11_satisfaction']:.1f}/7, {S['S5 Reluctant Instructors']['pct_dissatisfied']:.0f}% dissatisfied; publish/unpublish and Files-vs-Modules-vs-Pages confusion; Inbox messages unseen for weeks.",
   "I wanted to upload the slides after the lecture. It asked whether to put them in Files or Modules or Pages. I still don't know the difference; the TA fixed it later.",
   "A 'simple mode' with the five tasks they do, help at the moment of doing, Inbox forwarded to e-mail, one obvious place for files."),
 "S6 Access-First Learners": (
   f"{top_roles('S6 Access-First Learners')}. {S['S6 Access-First Learners']['access']['I use a screen reader or screen magnification']:.0f}% screen reader/magnification, {S['S6 Access-First Learners']['access']['I have dyslexia or another reading-related difference']:.0f}% dyslexia, {S['S6 Access-First Learners']['access']['I rely on captions or transcripts for audio and video']:.0f}% captions. n = {S['S6 Access-First Learners']['n']}.",
   f"Daily laptop users with the longest visits ({S['S6 Access-First Learners']['session']['More than 30 minutes']} of {S['S6 Access-First Learners']['n']} over 30 min) because every page costs effort; use Modules ({S['S6 Access-First Learners']['features']['Modules or Files (course materials)']:.0f}%), calendar ({S['S6 Access-First Learners']['features']['Calendar / To-do list']:.0f}%) and recordings; {S['S6 Access-First Learners']['tools']['Notion, Obsidian or OneNote']:.0f}% re-create course content in their own tools.",
   f"'Easy to read, see and hear' {S['S6 Access-First Learners']['likert_mean']['C9_accessible']:.1f}/7 vs {A['overall']['likert_mean']['C9_accessible']:.1f} overall; findability {S['S6 Access-First Learners']['likert_mean']['C1_find_fast']:.1f}/7; {S['S6 Access-First Learners']['pct_lost_2plus']:.0f}% failed to find something 2+ times; satisfaction {S['S6 Access-First Learners']['likert_mean']['C11_satisfaction']:.1f}/7, the lowest of all segments.",
   "The Modules page reads like a list of 200 links to my screen reader; the headings are not marked as headings so I cannot jump between weeks.",
   "Structure the system enforces (headings, alt text, captions, real text not scans), user-level display settings, full keyboard operation."),
}
for s in SEGS:
    who, beh, pains, quote, impl = DESC[s]
    cellw(wsg, r, 2, s, True, fill_hex=LAV); cellw(wsg, r, 3, who); cellw(wsg, r, 4, beh); cellw(wsg, r, 5, pains); cellw(wsg, r, 6, f'"{quote}"'); cellw(wsg, r, 7, impl)
    wsg.merge_cells(start_row=r, start_column=7, end_row=r, end_column=9); wsg.row_dimensions[r].height = 150; r += 1
for c in "CDEF": wsg.column_dimensions[c].width = 30
r += 1
# ---- PCA map data + scatter
section_title(wsg, r, "PCA map of respondents (first two principal components of the clustering features; chart data pasted)", 9); r += 1
note(wsg, f"B{r}", f"PC1 explains {A['pca_explained'][0]*100:.0f}% and PC2 {A['pca_explained'][1]*100:.0f}% of the variance. Each dot is a respondent, coloured by segment; overlap is expected in survey data (silhouette ~0.18)."); r += 1
pc_top = r
header_row(wsg, r, 2, [], fill_hex=GREY_D)
c0 = 2
colours = ["1B7F5E", "2563EB", "B45309", "7C3AED", "DC2626", "0891B2"]
sc = ScatterChart(); sc.title = "Respondents on PC1 / PC2 by segment"; sc.style = 13; sc.x_axis.title = "PC1"; sc.y_axis.title = "PC2"; sc.width, sc.height = 18, 12
sc.x_axis.delete = False; sc.y_axis.delete = False
for j, s in enumerate(SEGS):
    pts = [(float(x["PC1"]), float(x["PC2"])) for x in rows if x["Segment"] == s]
    cx, cy = c0 + 2 * j, c0 + 2 * j + 1
    cellw(wsg, r, cx, f"{s} PC1", True, fill_hex=GREY_D, size=8); cellw(wsg, r, cy, "PC2", True, fill_hex=GREY_D, size=8)
    for i, (x, y) in enumerate(pts, start=1):
        wsg.cell(row=r + i, column=cx, value=x).font = f(8); wsg.cell(row=r + i, column=cy, value=y).font = f(8)
    ser = Series(Reference(wsg, min_col=cy, min_row=r + 1, max_row=r + len(pts)), Reference(wsg, min_col=cx, min_row=r + 1, max_row=r + len(pts)), title=s)
    ser.marker.symbol = "circle"; ser.marker.size = 6; ser.graphicalProperties.line.noFill = True
    ser.marker.graphicalProperties.solidFill = colours[j]; ser.marker.graphicalProperties.line.solidFill = colours[j]
    sc.series.append(ser)
wsg.add_chart(sc, f"K{pc_top - 2}")
wsg.freeze_panes = "A4"

# =================================================================== 3-Persona
wq = wb.create_sheet("3-Persona")
P = S["S1 Deadline Sprinters"]
for k, v in {"A": 2, "B": 15, "C": 15, "D": 17, "E": 17, "F": 19, "G": 15, "H": 17, "I": 15, "J": 19, "K": 2}.items(): wq.column_dimensions[k].width = v
title(wq, "B1", "Step 4. Write a persona based on a segment  |  Persona 1 of 6, from S1 Deadline Sprinters (largest segment)", 13)
note(wq, "B2", "Layout follows the persona template from the Week 03 slides. Every statement in the card is backed by a row in the evidence table below (live formulas on the Responses sheet).", False); wq.merge_cells("B2:J2")
# block fills
def block(rng_, hex_, merge=True):
    if merge: wq.merge_cells(rng_)
    for row in wq[rng_]:
        for c in row: c.fill = fill(hex_); c.alignment = Alignment(wrap_text=True, vertical="top")
for rr in range(4, 24): wq.row_dimensions[rr].height = 22
# photo placeholder (B4:C11), name block (B12:C23)
block("B4:C11", PINK); wq["B4"] = "[ photo ]\n\nAdd an AI-generated face here, e.g. from thispersondoesnotexist.com (slide: Persona). Do not use a real student's photo."; wq["B4"].font = f(9, False, "6B7280", True); wq["B4"].alignment = CENTER
block("B12:C23", LAV_D, merge=False); wq.merge_cells("B12:C12"); wq.merge_cells("B13:C13")
wq["B12"] = "Sanne de Vries"; wq["B12"].font = f(16, True, "1E1B4B")
wq["B13"] = "The Deadline Sprinter"; wq["B13"].font = f(11, False, "3730A3", True)
for i, (icon, lab, val) in enumerate([("Age / gender", "Age / identifying gender", "21 / Female"), ("Location", "Location", "Student house in Tilburg (grew up in Eindhoven)"),
                                        ("Occupation", "Occupation", "2nd-year BSc Cognitive Science & AI; 10 h/week at a supermarket"), ("Family", "Family status", "Single, four housemates, no kids")]):
    rr = 15 + i * 2
    wq.cell(row=rr, column=2, value=lab).font = f(8, False, "4B5563"); wq.cell(row=rr + 1, column=2, value=val).font = f(9, True, "1E1B4B")
    wq.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3); wq.merge_cells(start_row=rr + 1, start_column=2, end_row=rr + 1, end_column=3)
    for cc in (2, 3): wq.cell(row=rr, column=cc).fill = fill(LAV_D); wq.cell(row=rr + 1, column=cc).fill = fill(LAV_D)
    wq.cell(row=rr + 1, column=2).alignment = WRAP
# quote (D4:E11), bio (D12:E23)
block("D4:E11", LAV)
wq["D4"] = ('"If it is not in my To-do list or in the WhatsApp group, it does not exist. I check Canvas on the bus, in the queue, in bed - '
            'but I only ever have thirty seconds."'); wq["D4"].font = f(10, True, GREEN, True)
block("D12:E23", LAV, merge=False); wq.merge_cells("D12:E12")
wq["D12"] = "Bio"; wq["D12"].font = f(12, True, GREEN)
wq["D13"] = (f"Sanne is in her second year of CSAI and takes four courses this block. She works two supermarket shifts a week and spends the rest of her time between lectures, the library and her friends. "
             f"Her phone is her main tool for everything, Canvas included: she opens the app several times a day for well under five minutes to see whether a grade came in, whether anything changed and what is due. "
             f"She has no patience for Modules on a phone screen; the slides she needs are forwarded in the course WhatsApp group anyway. She gets more Canvas notifications than she can read, so a change of deadline reaches her through a friend, not through Canvas, "
             f"and when a grade shows up without the class average she does not know what it means. She feels confident with software and expects apps to just work.")
wq["D13"].font = f(9); wq.merge_cells("D13:E23"); wq["D13"].alignment = WRAP
# goals (F4:G11), motivations (F12:G23)
block("F4:G11", PINK, merge=False); wq.merge_cells("F4:G4"); wq["F4"] = "Goals"; wq["F4"].font = f(12, True, GREEN)
wq["F5"] = ("- See in one glance what is due this week, across all courses\n- Know the moment a grade is published, with context (average, feedback)\n- Submit from the phone without hunting for the file\n- Never again travel to a lecture that was moved or cancelled\n- Spend no more than a minute per check")
wq["F5"].font = f(9); wq.merge_cells("F5:G11"); wq["F5"].alignment = WRAP
block("F12:G23", PINK, merge=False); wq.merge_cells("F12:G12"); wq.merge_cells("F13:G13"); wq["F12"] = "Motivations"; wq["F12"].font = f(12, True, GREEN); wq["F13"] = "(for managing her courses; share of segment, live)"; wq["F13"].font = f(8, False, "6B7280", True)
MOTIV = [("Not missing a deadline", f'=COUNTIFS({R("D1_rank__deadlines")},"<=2",{SEGRNG},"S1 Deadline Sprinters")/{SEGCELL["S1 Deadline Sprinters"]}'),
         ("Knowing where she stands (grades)", f'=COUNTIFS({R("B5_features__grades")},1,{SEGRNG},"S1 Deadline Sprinters")/{SEGCELL["S1 Deadline Sprinters"]}'),
         ("Speed: check in under 5 minutes", f'=(COUNTIFS({R("B4_session")},"lt2",{SEGRNG},"S1 Deadline Sprinters")+COUNTIFS({R("B4_session")},"2-5",{SEGRNG},"S1 Deadline Sprinters"))/{SEGCELL["S1 Deadline Sprinters"]}'),
         ("Staying in sync with friends", f'=COUNTIFS({R("B9_tools__chat")},1,{SEGRNG},"S1 Deadline Sprinters")/{SEGCELL["S1 Deadline Sprinters"]}'),
         ("Deep study of materials", f'=COUNTIFS({R("B5_features__recordings")},1,{SEGRNG},"S1 Deadline Sprinters")/{SEGCELL["S1 Deadline Sprinters"]}')]
for i, (lab, fm) in enumerate(MOTIV):
    rr = 14 + i * 2
    wq.cell(row=rr, column=6, value=lab).font = f(8, False, "374151"); wq.merge_cells(start_row=rr, start_column=6, end_row=rr, end_column=7)
    wq.cell(row=rr + 1, column=6, value=f'=REPT("|",ROUND(G{rr+1}*30,0))').font = Font(name="Arial", size=8, color="F59E0B", bold=True)
    wq.cell(row=rr + 1, column=7, value=fm).number_format = "0%"; wq.cell(row=rr + 1, column=7).font = f(9, True, "1E1B4B"); wq.cell(row=rr + 1, column=7).alignment = Alignment(horizontal="right")
    for cc in (6, 7): wq.cell(row=rr, column=cc).fill = fill(PINK); wq.cell(row=rr + 1, column=cc).fill = fill(PINK)
# pains (H4:I11), devices (H12:I23)
block("H4:I11", PINK, merge=False); wq.merge_cells("H4:I4"); wq["H4"] = "Pains"; wq["H4"].font = f(12, True, GREEN)
wq["H5"] = (f"- Too many irrelevant notifications ({P['likert_mean']['C4_notif_noise']:.1f}/7), so the one that matters gets lost\n- Canvas on the phone is worse than on a laptop ({P['likert_mean']['C5_mobile']:.1f}/7): re-logins, feedback invisible, file picker fails\n"
            f"- Every course puts deadlines and slides somewhere else ({P['likert_mean']['C6_consistent']:.1f}/7)\n- To-do list polluted by undated and old items\n- Grades without context; 'submitted' for days without knowing the file was wrong")
wq["H5"].font = f(9); wq.merge_cells("H5:I11"); wq["H5"].alignment = WRAP
block("H12:I23", PINK, merge=False); wq.merge_cells("H12:I12"); wq.merge_cells("H13:I13"); wq["H12"] = "Devices"; wq["H12"].font = f(12, True, GREEN); wq["H13"] = "(most-used device in the segment, live)"; wq["H13"].font = f(8, False, "6B7280", True)
DEV = [("Smartphone (Canvas app)", '"phone"'), ("Laptop", '"laptop"'), ("Tablet", '"tablet"'), ("Desktop", '"desktop"')]
for i, (lab, key) in enumerate(DEV):
    rr = 14 + i * 2
    wq.cell(row=rr, column=8, value=lab).font = f(8, False, "374151"); wq.merge_cells(start_row=rr, start_column=8, end_row=rr, end_column=9)
    wq.cell(row=rr + 1, column=8, value=f'=REPT("|",ROUND(I{rr+1}*30,0))').font = Font(name="Arial", size=8, color="2563EB", bold=True)
    wq.cell(row=rr + 1, column=9, value=f'=COUNTIFS({R("B3_primary")},{key},{SEGRNG},"S1 Deadline Sprinters")/{SEGCELL["S1 Deadline Sprinters"]}').number_format = "0%"
    wq.cell(row=rr + 1, column=9).font = f(9, True, "1E1B4B"); wq.cell(row=rr + 1, column=9).alignment = Alignment(horizontal="right")
    for cc in (8, 9): wq.cell(row=rr, column=cc).fill = fill(PINK); wq.cell(row=rr + 1, column=cc).fill = fill(PINK)
wq["H22"] = "Uses the laptop for submitting long assignments and for the rare deep session; never a desktop."; wq["H22"].font = f(8, False, "6B7280", True); wq.merge_cells("H22:I23"); wq["H22"].alignment = WRAP
# brands (J4:J23)
block("J4:J23", PINK, merge=False); wq["J4"] = "Brand affiliations"; wq["J4"].font = f(12, True, GREEN)
wq["J6"] = ("WhatsApp (course groups)\n\nCanvas Student app\n\nGoogle Calendar\n\nInstagram / TikTok\n\nSpotify\n\nNS app (train)\n\nAlbert Heijn (side job)\n\nBol / Studystore (books)")
wq["J6"].font = f(9); wq.merge_cells("J6:J23"); wq["J6"].alignment = WRAP
# ---- evidence table
r = 26
section_title(wq, r, "Evidence: every statement in the card traces back to the S1 segment data (live formulas)", 10); r += 1
header_row(wq, r, 2, ["Persona statement", "", "", "Statistic in S1 Deadline Sprinters", "", "Value", "Same statistic, all respondents", "", "Source item"], fill_hex=GREY_D, color="222222")
wq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); wq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6); wq.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9); r += 1
s1 = "S1 Deadline Sprinters"; n1 = SEGCELL[s1]
def ev(statement, stat, seg_formula, all_formula, fmt, src):
    global r
    cellw(wq, r, 2, statement); wq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cellw(wq, r, 5, stat); wq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    cellw(wq, r, 7, seg_formula, True, fmt=fmt, align=CENTER, fill_hex=LAV)
    cellw(wq, r, 8, all_formula, fmt=fmt, align=CENTER); wq.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    cellw(wq, r, 10, src, align=CENTER); wq.row_dimensions[r].height = 30; r += 1
ev("21, female, Bachelor student", "% Bachelor students / % aged 17-24 / % female", f'=COUNTIFS({R("A1_role")},"bsc",{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("A1_role")},"bsc")/{NCELL}', "0%", "A1")
ev("", "% aged 17-24", f'=(COUNTIFS({R("A2_age")},"17-20",{SEGRNG},"{s1}")+COUNTIFS({R("A2_age")},"21-24",{SEGRNG},"{s1}"))/{n1}', f'=(COUNTIF({R("A2_age")},"17-20")+COUNTIF({R("A2_age")},"21-24"))/{NCELL}', "0%", "A2")
ev("", "% female", f'=COUNTIFS({R("A3_gender")},"female",{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("A3_gender")},"female")/{NCELL}', "0%", "A3")
ev("Works ~10 h/week alongside her studies", "% with 1-16 h/week outside commitments", f'=(COUNTIFS({R("A7_hours")},"1-8",{SEGRNG},"{s1}")+COUNTIFS({R("A7_hours")},"9-16",{SEGRNG},"{s1}"))/{n1}', f'=(COUNTIF({R("A7_hours")},"1-8")+COUNTIF({R("A7_hours")},"9-16"))/{NCELL}', "0%", "A7")
ev("Takes four courses", "Mean number of courses", f'=AVERAGEIFS({R("courses_v")},{SEGRNG},"{s1}")', f'=AVERAGE({R("courses_v")})', "0.0", "A6")
ev("Opens the app several times a day", "% opening at least once a day", f'=(COUNTIFS({R("B1_freq")},"daily",{SEGRNG},"{s1}")+COUNTIFS({R("B1_freq")},"several",{SEGRNG},"{s1}"))/{n1}', f'=(COUNTIF({R("B1_freq")},"daily")+COUNTIF({R("B1_freq")},"several"))/{NCELL}', "0%", "B1")
ev("Her phone is her main tool", "% smartphone as most-used device", f'=COUNTIFS({R("B3_primary")},"phone",{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B3_primary")},"phone")/{NCELL}', "0%", "B3")
ev("", "% used the Canvas app in the last 7 days", f'=COUNTIFS({R("B2_devices__phone_app")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B2_devices__phone_app")},1)/{NCELL}', "0%", "B2")
ev("Visits last well under five minutes", "% with typical visit under 5 minutes", f'=(COUNTIFS({R("B4_session")},"lt2",{SEGRNG},"{s1}")+COUNTIFS({R("B4_session")},"2-5",{SEGRNG},"{s1}"))/{n1}', f'=(COUNTIF({R("B4_session")},"lt2")+COUNTIF({R("B4_session")},"2-5"))/{NCELL}', "0%", "B4")
ev("Checks whether a grade came in", "% used Grades/feedback last week", f'=COUNTIFS({R("B5_features__grades")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B5_features__grades")},1)/{NCELL}', "0%", "B5")
ev("Has no patience for Modules on a phone", "% used Modules/Files last week", f'=COUNTIFS({R("B5_features__modules")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B5_features__modules")},1)/{NCELL}', "0%", "B5")
ev("Deadlines and grades are what matter", "Mean rank of 'keeping track of deadlines' (1 = top)", f'=AVERAGEIFS({R("D1_rank__deadlines")},{SEGRNG},"{s1}")', f'=AVERAGE({R("D1_rank__deadlines")})', "0.0", "D1")
ev("", "Mean rank of 'seeing grades and feedback'", f'=AVERAGEIFS({R("D1_rank__grades")},{SEGRNG},"{s1}")', f'=AVERAGE({R("D1_rank__grades")})', "0.0", "D1")
ev("A deadline change reaches her through a friend", "% learned of last deadline from a classmate (e.g. WhatsApp)", f'=COUNTIFS({R("B6_deadline__peer")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B6_deadline__peer")},1)/{NCELL}', "0%", "B6")
ev("", "% learned of it from a Canvas push notification", f'=COUNTIFS({R("B6_deadline__push")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B6_deadline__push")},1)/{NCELL}', "0%", "B6")
ev("Lives in the course WhatsApp group", "% using a WhatsApp/Signal/Discord group for courses", f'=COUNTIFS({R("B9_tools__chat")},1,{SEGRNG},"{s1}")/{n1}', f'=COUNTIF({R("B9_tools__chat")},1)/{NCELL}', "0%", "B9")
ev("More notifications than she can read", "'I receive notifications that are not relevant to me' (1-7)", f'=AVERAGEIFS({R("C4_notif_noise")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C4_notif_noise")})', "0.0", "C4")
ev("", "'Notifications help me stay on top of what is due' (1-7)", f'=AVERAGEIFS({R("C3_notif_useful")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C3_notif_useful")})', "0.0", "C3")
ev("Canvas on the phone is worse than on a laptop", "'Using the system on my phone works as well as on a computer' (1-7)", f'=AVERAGEIFS({R("C5_mobile")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C5_mobile")})', "0.0", "C5")
ev("Grades without context, not sure they are current", "'I trust that grades and feedback are up to date' (1-7)", f'=AVERAGEIFS({R("C7_uptodate")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C7_uptodate")})', "0.0", "C7")
ev("Every course puts things somewhere else", "'Course pages are organised in a similar way' (1-7)", f'=AVERAGEIFS({R("C6_consistent")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C6_consistent")})', "0.0", "C6")
ev("Confident with software, expects apps to just work", "'I can complete my tasks without asking for help' (1-7) / comfort learning software (1-7)", f'=AVERAGEIFS({R("C8_confident")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C8_confident")})', "0.0", "C8")
ev("", "Comfort learning new software alone (1-7)", f'=AVERAGEIFS({R("A9_tech")},{SEGRNG},"{s1}")', f'=AVERAGE({R("A9_tech")})', "0.0", "A9")
ev("Moderately satisfied overall (it works for quick checks)", "Overall satisfaction (1-7)", f'=AVERAGEIFS({R("C11_satisfaction")},{SEGRNG},"{s1}")', f'=AVERAGE({R("C11_satisfaction")})', "0.0", "C11")
r += 1
# ---- POV and HMW
section_title(wq, r, "Point of view (Week 03 slides, 4-sentence formulation) and How-Might-We questions", 10); r += 1
POV = [("WE MET", "Sanne, a second-year CSAI student in Tilburg who opens the Canvas app on her phone several times a day, for thirty seconds at a time."),
       ("WE WERE SURPRISED TO NOTICE", f"that although she receives more Canvas notifications than anyone, she learns about deadline changes from the course WhatsApp group ({P['deadline']['A classmate or colleague told me (e.g. WhatsApp)']:.0f}% of her segment did for their last deadline)."),
       ("WE WONDER IF THIS MEANS", "that for her the official channel has become background noise and her peers have become the trusted filter: she outsources 'what matters' to people who also do not know."),
       ("IT WOULD BE GAME-CHANGING TO", "make the one change that matters unmistakable on a phone screen within five seconds, without her having to open a single course page.")]
for lab, txt in POV:
    cellw(wq, r, 2, lab, True, color=TEAL_D); wq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cellw(wq, r, 5, txt); wq.merge_cells(start_row=r, start_column=5, end_row=r, end_column=10); wq.row_dimensions[r].height = 32; r += 1
r += 1
cellw(wq, r, 2, "How might we ...", True, color=TEAL_D); wq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); r += 1
for h in ["... make a change to a deadline impossible to miss on a phone, in under five seconds?",
          "... show what is due this week across all courses even when lecturers have not dated their assignments?",
          "... give a grade its meaning the moment it appears (context, feedback, what to do next) on a small screen?",
          "... bring the trust Sanne places in her WhatsApp group into the official channel, instead of competing with it?",
          "... let a thirty-second check on the bus be enough, so that nothing requires the laptop later?"]:
    cellw(wq, r, 2, h); wq.merge_cells(start_row=r, start_column=2, end_row=r, end_column=10); r += 1
r += 1
# ---- persona candidates for other segments
section_title(wq, r, "Persona candidates for the other five segments (to be written the same way; n and key numbers are live)", 10); r += 1
header_row(wq, r, 2, ["Segment", "n", "Suggested persona", "", "Tagline", "", "Key numbers (live)", "", "Top pain"], fill_hex=GREY_D, color="222222")
wq.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5); wq.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7); wq.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9); r += 1
CAND = {
 "S2 Structured Scholars": ("Tomasz Nowak, 24, MSc Data Science (from Krakow), laptop + Notion", "The Structured Scholar: builds his own system because every course is different",
                             lambda s: f'="Consistency "&TEXT(AVERAGEIFS({R("C6_consistent")},{SEGRNG},"{s}"),"0.0")&"/7 · Modules "&TEXT(COUNTIFS({R("B5_features__modules")},1,{SEGRNG},"{s}")/{SEGCELL[s]},"0%")&" · Notes app "&TEXT(COUNTIFS({R("B9_tools__notes")},1,{SEGRNG},"{s}")/{SEGCELL[s]},"0%")',
                             "Course pages organised differently everywhere; no search across courses or inside files."),
 "S3 Overloaded Jugglers": ("Fatima el Amrani, 31, part-time BSc Psychology, 24 h/week job, one child", "The Overloaded Juggler: catches up on Sunday evening, by e-mail",
                             lambda s: f'="17+ h outside "&TEXT((COUNTIFS({R("A7_hours")},"17-24",{SEGRNG},"{s}")+COUNTIFS({R("A7_hours")},"25+",{SEGRNG},"{s}"))/{SEGCELL[s]},"0%")&" · missed deadline "&TEXT(COUNTIFS({R("B6_deadline__late")},1,{SEGRNG},"{s}")/{SEGCELL[s]},"0%")&" · freq "&TEXT(AVERAGEIFS({R("freq_v")},{SEGRNG},"{s}"),"0.0")&"/6"',
                             "No 'what did I miss / what must I do' view; deadline changes buried in look-alike e-mails."),
 "S4 Power Instructors": ("Dr. Eva Janssen, 38, assistant professor, 2 courses, 180 students, desktop all day", "The Power Instructor: fights the system with bulk actions and Excel",
                           lambda s: f'="Daily "&TEXT((COUNTIFS({R("B1_freq")},"daily",{SEGRNG},"{s}")+COUNTIFS({R("B1_freq")},"several",{SEGRNG},"{s}"))/{SEGCELL[s]},"0%")&" · notif. noise "&TEXT(AVERAGEIFS({R("C4_notif_noise")},{SEGRNG},"{s}"),"0.0")&"/7 · group support "&TEXT(AVERAGEIFS({R("C10_group")},{SEGRNG},"{s}"),"0.0")&"/7"',
                           "Repetitive clicks: course-copy dates, SpeedGrader reloads, section overrides, groups ignoring sections."),
 "S5 Reluctant Instructors": ("Prof. Henk van der Berg, 57, full professor, delegates Canvas to his TA", "The Reluctant Instructor: once a week, on the desktop, with the education office on speed dial",
                              lambda s: f'="Confidence "&TEXT(AVERAGEIFS({R("C8_confident")},{SEGRNG},"{s}"),"0.0")&"/7 · last task ease "&TEXT(AVERAGEIFS({R("C12_ease")},{SEGRNG},"{s}"),"0.0")&"/5 · e-mail as tool "&TEXT(COUNTIFS({R("B9_tools__email")},1,{SEGRNG},"{s}")/{SEGCELL[s]},"0%")',
                              "Files vs Modules vs Pages; invisible 'unpublished' state; Inbox messages unseen for weeks."),
 "S6 Access-First Learners": ("Lars Pedersen, 23, BSc Data Science, screen-reader and dyslexia support user", "The Access-First Learner: every page is work; structure is everything",
                              lambda s: f'="Readable "&TEXT(AVERAGEIFS({R("C9_accessible")},{SEGRNG},"{s}"),"0.0")&"/7 · findability "&TEXT(AVERAGEIFS({R("C1_find_fast")},{SEGRNG},"{s}"),"0.0")&"/7 · satisfaction "&TEXT(AVERAGEIFS({R("C11_satisfaction")},{SEGRNG},"{s}"),"0.0")&"/7"',
                              "Unstructured pages, scanned PDFs, auto-captions, colour-only status, keyboard traps."),
}
for s in SEGS[1:]:
    who, tag, keyf, pain = CAND[s]
    cellw(wq, r, 2, s, True, fill_hex=LAV); cellw(wq, r, 3, f'=COUNTIF({SEGRNG},"{s}")', align=CENTER)
    cellw(wq, r, 4, who); wq.merge_cells(start_row=r, start_column=4, end_row=r, end_column=5)
    cellw(wq, r, 6, tag); wq.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    cellw(wq, r, 8, keyf(s)); wq.merge_cells(start_row=r, start_column=8, end_row=r, end_column=9)
    cellw(wq, r, 10, pain); wq.row_dimensions[r].height = 58; r += 1
note(wq, f"B{r+1}", "Persona names and photos are fictional; all characteristics come from the segment statistics. Next step (Week 03): brainstorm 10-15 HMWs per POV, pick the best three, then storyboard.")
wq.freeze_panes = "A4"

# ------------------------------------------------------------------- finish
wb.active = 0
for w in wb.worksheets:
    w.sheet_view.showGridLines = (w.title in ("Responses", "Codes"))
wb.save("out/iCanvas_needfinding.xlsx")
print("saved out/iCanvas_needfinding.xlsx")
